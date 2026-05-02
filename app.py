import streamlit as st
import pandas as pd
import os
import io
import time
import csv
import json
import base64
import re
import hashlib
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from datetime import datetime
from google import genai
from pypdf import PdfReader

# ReportLab for PDF generation
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
)
from reportlab.lib.enums import TA_LEFT, TA_CENTER

# python-docx for Word export
try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# ─── Page Config ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="FAQ Forge",
    page_icon="◈",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ─── Custom CSS ───────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=Fraunces:ital,opsz,wght@0,9..144,300;0,9..144,600;0,9..144,700;1,9..144,400&family=JetBrains+Mono:wght@400;500&display=swap');

:root {
    --bg:        #0e0e14;
    --bg2:       #13131c;
    --surface:   #16161f;
    --surface2:  #1c1c28;
    --surface3:  #222233;
    --border:    rgba(138,110,255,0.14);
    --border-hi: rgba(138,110,255,0.35);

    --violet:    #8a6eff;
    --violet2:   #6c4ef0;
    --cyan:      #22d3ee;
    --cyan2:     #06b6d4;
    --pink:      #f472b6;
    --green:     #34d399;
    --amber:     #fbbf24;

    --text:      #f0eeff;
    --text2:     #c4bfdf;
    --muted:     #7c769e;
    --dim:       #4a4568;
}

/* ── Light mode overrides ── */
body.light-mode {
    --bg:        #f4f2ff;
    --bg2:       #ebe8ff;
    --surface:   #ffffff;
    --surface2:  #f0eeff;
    --surface3:  #e4dfff;
    --border:    rgba(108,78,240,0.15);
    --border-hi: rgba(108,78,240,0.35);
    --text:      #1a1035;
    --text2:     #3d3460;
    --muted:     #6b5ea8;
    --dim:       #9b90c8;
}

/* ── Theme toggle button ── */
.theme-toggle {
    position: fixed;
    top: 1rem;
    right: 5rem;
    z-index: 9999;
    background: var(--surface2);
    border: 1px solid var(--border-hi);
    border-radius: 50px;
    padding: 0.4rem 1rem;
    cursor: pointer;
    font-size: 0.8rem;
    font-family: 'Plus Jakarta Sans', sans-serif;
    font-weight: 600;
    color: var(--text);
    transition: all 0.2s;
}
.theme-toggle:hover {
    background: var(--violet2);
    color: white;
}

* { box-sizing: border-box; }

html, body, [class*="css"] {
    font-family: 'Plus Jakarta Sans', sans-serif;
    color: var(--text);
}

/* ── App background with grid ── */
.stApp {
    background-color: var(--bg);
    background-image:
        linear-gradient(rgba(138,110,255,0.03) 1px, transparent 1px),
        linear-gradient(90deg, rgba(138,110,255,0.03) 1px, transparent 1px),
        radial-gradient(ellipse 70% 50% at 15% 0%, rgba(138,110,255,0.08) 0%, transparent 55%),
        radial-gradient(ellipse 50% 40% at 85% 100%, rgba(34,211,238,0.06) 0%, transparent 55%);
    background-size: 40px 40px, 40px 40px, 100% 100%, 100% 100%;
    min-height: 100vh;
}

/* ── Main container ── */
.main .block-container {
    padding: 3rem 3.5rem 4rem;
    max-width: 1200px;
}

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background: var(--bg2) !important;
    border-right: 1px solid var(--border) !important;
}
[data-testid="stSidebar"] .block-container {
    padding: 2.5rem 1.6rem !important;
}
[data-testid="stSidebar"] * { color: var(--muted) !important; }
[data-testid="stSidebar"] h3 {
    color: var(--text2) !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.65rem !important;
    font-weight: 700 !important;
    letter-spacing: 0.18em !important;
    text-transform: uppercase !important;
    margin-bottom: 1rem !important;
    padding-bottom: 0.6rem !important;
    border-bottom: 1px solid var(--border) !important;
}
[data-testid="stSidebar"] .stTextInput input {
    background: var(--surface) !important;
    border: 1px solid var(--border-hi) !important;
    border-radius: 8px !important;
    color: var(--text) !important;
    font-family: 'JetBrains Mono', monospace !important;
    font-size: 0.78rem !important;
}

/* ── Page header ── */
.page-header {
    display: flex;
    align-items: center;
    gap: 1.4rem;
    margin-bottom: 3rem;
    padding-bottom: 2rem;
    border-bottom: 1px solid var(--border);
}
.header-icon {
    width: 56px;
    height: 56px;
    background: linear-gradient(135deg, var(--violet2), var(--cyan2));
    border-radius: 14px;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 1.6rem;
    flex-shrink: 0;
    box-shadow: 0 0 28px rgba(138,110,255,0.35), 0 0 60px rgba(34,211,238,0.12);
}
.header-title {
    font-family: 'Fraunces', serif;
    font-size: 3rem;
    font-weight: 700;
    letter-spacing: -0.02em;
    line-height: 1.2;
    margin: 0;
    padding-bottom: 4px;
    background: linear-gradient(100deg, #f0eeff 20%, #8a6eff 55%, #22d3ee 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
}
.header-sub {
    font-size: 0.8rem;
    color: var(--muted);
    font-weight: 500;
    letter-spacing: 0.06em;
    text-transform: uppercase;
    margin-top: 0.4rem;
}

/* ── Form labels ── */
.stTextArea label, .stTextInput label,
.stSelectbox label, .stSlider label {
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.68rem !important;
    font-weight: 700 !important;
    letter-spacing: 0.15em !important;
    text-transform: uppercase !important;
    color: var(--muted) !important;
}

/* ── Textarea ── */
.stTextArea textarea {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 12px !important;
    color: var(--text) !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.92rem !important;
    line-height: 1.75 !important;
    transition: border-color 0.2s, box-shadow 0.2s !important;
}
.stTextArea textarea:focus {
    border-color: var(--violet) !important;
    box-shadow: 0 0 0 3px rgba(138,110,255,0.12) !important;
}

/* ── Text input ── */
.stTextInput input {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px !important;
    color: var(--text) !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.92rem !important;
    transition: border-color 0.2s !important;
}
.stTextInput input:focus {
    border-color: var(--violet) !important;
    box-shadow: 0 0 0 3px rgba(138,110,255,0.12) !important;
}

/* ── Selectbox ── */
.stSelectbox > div > div {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px !important;
    color: var(--text) !important;
}

/* ── Download buttons ── */
.stDownloadButton > button {
    background: var(--surface2) !important;
    color: var(--cyan) !important;
    border: 1px solid rgba(34,211,238,0.3) !important;
    border-radius: 10px !important;
    padding: 0.65rem 1.2rem !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.75rem !important;
    letter-spacing: 0.1em !important;
    text-transform: uppercase !important;
    transition: all 0.2s ease !important;
    width: 100%;
}
.stDownloadButton > button:hover {
    background: rgba(34,211,238,0.1) !important;
    border-color: var(--cyan) !important;
    box-shadow: 0 0 18px rgba(34,211,238,0.2) !important;
    transform: translateY(-1px) !important;
}

/* ── Primary / Generate button ── */
.stButton > button {
    background: linear-gradient(135deg, var(--violet2) 0%, #5b3fd8 50%, var(--cyan2) 100%) !important;
    color: #fff !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 0.8rem 1.5rem !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.85rem !important;
    letter-spacing: 0.08em !important;
    text-transform: uppercase !important;
    transition: all 0.25s ease !important;
    width: 100%;
    box-shadow: 0 4px 20px rgba(108,78,240,0.45) !important;
    position: relative;
}
.stButton > button:hover {
    transform: translateY(-2px) !important;
    box-shadow: 0 8px 30px rgba(108,78,240,0.55), 0 0 40px rgba(34,211,238,0.15) !important;
    filter: brightness(1.08) !important;
}
.stButton > button:active {
    transform: translateY(0) !important;
}

/* Clear button */
.clear-btn .stButton > button {
    background: transparent !important;
    color: var(--muted) !important;
    border: 1px solid var(--border) !important;
    font-size: 0.72rem !important;
    box-shadow: none !important;
    filter: none !important;
}
.clear-btn .stButton > button:hover {
    border-color: rgba(244,114,182,0.5) !important;
    color: var(--pink) !important;
    background: rgba(244,114,182,0.05) !important;
    box-shadow: none !important;
    transform: none !important;
}

/* ── Config panel ── */
.config-panel {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 16px;
    padding: 1.8rem;
    box-shadow: 0 4px 24px rgba(0,0,0,0.3);
}

/* ── FAQ result ── */
.faq-result-wrapper { margin-top: 2.5rem; }
.faq-result-header {
    display: flex;
    align-items: center;
    justify-content: space-between;
    margin-bottom: 1rem;
    padding-bottom: 1rem;
    border-bottom: 1px solid var(--border);
}
.faq-result-title {
    font-family: 'Fraunces', serif;
    font-size: 1.65rem;
    font-weight: 600;
    color: var(--text);
    letter-spacing: -0.01em;
}
.faq-tag {
    display: inline-block;
    background: var(--surface2);
    border: 1px solid var(--border);
    color: var(--muted);
    font-size: 0.65rem;
    font-weight: 700;
    letter-spacing: 0.12em;
    text-transform: uppercase;
    border-radius: 100px;
    padding: 0.25rem 0.75rem;
}
.faq-block {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 14px;
    overflow: hidden;
    box-shadow: 0 4px 24px rgba(0,0,0,0.25);
}
.faq-item {
    padding: 1.5rem 1.8rem;
    border-bottom: 1px solid var(--border);
    transition: background 0.15s;
    position: relative;
}
.faq-item::before {
    content: '';
    position: absolute;
    left: 0; top: 0; bottom: 0;
    width: 3px;
    background: transparent;
    transition: background 0.2s;
    border-radius: 0 2px 2px 0;
}
.faq-item:hover { background: var(--surface2); }
.faq-item:hover::before {
    background: linear-gradient(180deg, var(--violet), var(--cyan));
}
.faq-item:last-child { border-bottom: none; }
.faq-q {
    font-family: 'Plus Jakarta Sans', sans-serif;
    font-size: 0.97rem;
    font-weight: 700;
    color: var(--text);
    margin-bottom: 0.65rem;
    display: flex;
    gap: 0.85rem;
    align-items: flex-start;
    line-height: 1.5;
}
.faq-q-num {
    font-family: 'JetBrains Mono', monospace;
    font-size: 0.65rem;
    color: var(--violet);
    background: rgba(138,110,255,0.12);
    border: 1px solid rgba(138,110,255,0.25);
    border-radius: 6px;
    padding: 0.22rem 0.55rem;
    margin-top: 0.15rem;
    white-space: nowrap;
    flex-shrink: 0;
    letter-spacing: 0.06em;
}
.faq-a {
    font-size: 0.875rem;
    color: var(--text2);
    line-height: 1.8;
    padding-left: 2.7rem;
    font-weight: 400;
}

/* Raw fallback */
.faq-raw {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 1.8rem 2rem;
    color: var(--text2);
    font-size: 0.875rem;
    line-height: 1.85;
    white-space: pre-wrap;
    font-family: 'JetBrains Mono', monospace;
}

/* ── Divider ── */
hr {
    border: none !important;
    border-top: 1px solid var(--border) !important;
    margin: 2.5rem 0 !important;
}

/* ── Subheaders ── */
h2, h3 {
    font-family: 'Fraunces', serif !important;
    font-weight: 600 !important;
    color: var(--text) !important;
    letter-spacing: -0.01em !important;
}

/* ── History rows ── */
.history-row {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 1rem 1.4rem;
    margin-bottom: 0.55rem;
    display: flex;
    align-items: center;
    justify-content: space-between;
    transition: border-color 0.15s, box-shadow 0.15s;
}
.history-row:hover {
    border-color: var(--border-hi);
    box-shadow: 0 2px 12px rgba(138,110,255,0.1);
}
.history-topic {
    font-weight: 700;
    font-size: 0.9rem;
    color: var(--text);
}
.history-meta {
    font-size: 0.68rem;
    color: var(--muted);
    margin-top: 0.15rem;
    font-family: 'JetBrains Mono', monospace;
    letter-spacing: 0.04em;
}

/* Difficulty badges */
.badge-easy   { background: rgba(52,211,153,0.12); color: #34d399; border: 1px solid rgba(52,211,153,0.25); }
.badge-medium { background: rgba(251,191,36,0.12); color: #fbbf24; border: 1px solid rgba(251,191,36,0.25); }
.badge-hard   { background: rgba(244,114,182,0.12); color: #f472b6; border: 1px solid rgba(244,114,182,0.25); }
.diff-badge {
    display: inline-block;
    font-size: 0.65rem;
    font-weight: 700;
    letter-spacing: 0.12em;
    text-transform: uppercase;
    border-radius: 100px;
    padding: 0.22rem 0.8rem;
}

/* ── Expander ── */
.streamlit-expanderHeader {
    background: var(--surface2) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px !important;
    color: var(--text) !important;
}

/* ── Alerts ── */
.stSuccess {
    background: rgba(52,211,153,0.07) !important;
    border: 1px solid rgba(52,211,153,0.22) !important;
    border-left: 3px solid #34d399 !important;
    border-radius: 10px !important;
    color: #6ee7b7 !important;
}
.stInfo {
    background: rgba(138,110,255,0.07) !important;
    border: 1px solid rgba(138,110,255,0.22) !important;
    border-left: 3px solid var(--violet) !important;
    border-radius: 10px !important;
}
.stWarning {
    background: rgba(251,191,36,0.07) !important;
    border: 1px solid rgba(251,191,36,0.22) !important;
    border-left: 3px solid var(--amber) !important;
    border-radius: 10px !important;
}
.stError {
    background: rgba(244,114,182,0.07) !important;
    border: 1px solid rgba(244,114,182,0.22) !important;
    border-left: 3px solid var(--pink) !important;
    border-radius: 10px !important;
}

/* ── Spinner ── */
.stSpinner > div { border-top-color: var(--violet) !important; }

/* ── Dataframe ── */
.stDataFrame {
    border-radius: 10px;
    overflow: hidden;
    border: 1px solid var(--border) !important;
}

/* ── Input mode tabs ── */
.stTabs [data-baseweb="tab-list"] {
    background: var(--surface2) !important;
    border-radius: 10px !important;
    padding: 4px !important;
    gap: 2px !important;
    border: 1px solid var(--border) !important;
}
.stTabs [data-baseweb="tab"] {
    background: transparent !important;
    border-radius: 8px !important;
    color: var(--muted) !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.78rem !important;
    font-weight: 600 !important;
    letter-spacing: 0.06em !important;
    padding: 0.5rem 1.2rem !important;
    transition: all 0.2s !important;
}
.stTabs [aria-selected="true"] {
    background: linear-gradient(135deg, var(--violet2), var(--cyan2)) !important;
    color: white !important;
}
.stTabs [data-baseweb="tab-panel"] {
    padding-top: 1.2rem !important;
}

/* ── File uploader ── */
.stFileUploader {
    background: var(--surface) !important;
    border: 2px dashed var(--border-hi) !important;
    border-radius: 12px !important;
    padding: 1rem !important;
}
.stFileUploader:hover {
    border-color: var(--violet) !important;
}
[data-testid="stFileUploaderDropzone"] {
    background: var(--surface) !important;
}

/* ── Extracted content preview ── */
.extracted-preview {
    background: var(--surface2);
    border: 1px solid var(--border);
    border-left: 3px solid var(--violet);
    border-radius: 0 10px 10px 0;
    padding: 1rem 1.2rem;
    margin-top: 0.8rem;
    font-size: 0.8rem;
    color: var(--text2);
    line-height: 1.7;
    max-height: 180px;
    overflow-y: auto;
    font-family: 'JetBrains Mono', monospace;
    white-space: pre-wrap;
}
.extracted-label {
    font-size: 0.65rem;
    font-weight: 700;
    letter-spacing: 0.15em;
    text-transform: uppercase;
    color: var(--violet);
    margin-bottom: 0.5rem;
}
/* ── Login page ── */
.login-wrapper {
    max-width: 420px;
    margin: 6rem auto;
    background: var(--surface);
    border: 1px solid var(--border-hi);
    border-radius: 20px;
    padding: 2.8rem 2.5rem;
    box-shadow: 0 8px 40px rgba(108,78,240,0.2);
}
.login-logo {
    text-align: center;
    font-size: 2.5rem;
    margin-bottom: 0.5rem;
}
.login-title {
    font-family: 'Fraunces', serif;
    font-size: 1.8rem;
    font-weight: 700;
    text-align: center;
    background: linear-gradient(100deg, #f0eeff 20%, #8a6eff 55%, #22d3ee 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    margin-bottom: 0.3rem;
}
.login-sub {
    text-align: center;
    font-size: 0.78rem;
    color: var(--muted);
    margin-bottom: 2rem;
    letter-spacing: 0.04em;
}
.user-badge {
    display: inline-flex;
    align-items: center;
    gap: 0.5rem;
    background: rgba(138,110,255,0.10);
    border: 1px solid rgba(138,110,255,0.25);
    border-radius: 100px;
    padding: 0.3rem 0.9rem;
    font-size: 0.75rem;
    font-weight: 600;
    color: var(--violet);
}
/* ── Rating buttons ── */
.rating-row {
    display: flex;
    gap: 0.5rem;
    margin-top: 0.4rem;
    padding-left: 2.7rem;
}
.rate-btn {
    background: var(--surface2);
    border: 1px solid var(--border);
    border-radius: 6px;
    padding: 0.2rem 0.7rem;
    font-size: 0.75rem;
    cursor: pointer;
    transition: all 0.15s;
    color: var(--muted);
}
.rate-btn:hover { border-color: var(--violet); color: var(--violet); }
.rate-up   { color: #34d399 !important; border-color: rgba(52,211,153,0.3) !important; }
.rate-down { color: #f472b6 !important; border-color: rgba(244,114,182,0.3) !important; }
</style>
""", unsafe_allow_html=True)

# ─── Helpers ─────────────────────────────────────────────────────────────────

# ── User database ─────────────────────────────────────────────────────────────
USERS_FILE = "faq_users.json"

def _hash(password: str) -> str:
    return hashlib.sha256(password.encode()).hexdigest()

def load_users() -> dict:
    if os.path.exists(USERS_FILE):
        with open(USERS_FILE, "r") as f:
            return json.load(f)
    default = {"admin": {"password": _hash("admin123"), "created": str(datetime.now().date())}}
    save_users(default)
    return default

def save_users(users: dict):
    with open(USERS_FILE, "w") as f:
        json.dump(users, f, indent=2)

def register_user(username: str, password: str) -> bool:
    users = load_users()
    if username in users:
        return False
    users[username] = {"password": _hash(password), "created": str(datetime.now().date())}
    save_users(users)
    return True

def verify_user(username: str, password: str) -> bool:
    users = load_users()
    return username in users and users[username]["password"] == _hash(password)

# ── User-aware history ────────────────────────────────────────────────────────
def history_file(username: str = "guest") -> str:
    return f"faq_history_{username}.csv"

def load_history(username: str = "guest") -> pd.DataFrame:
    f = history_file(username)
    if os.path.exists(f):
        df = pd.read_csv(f)
        if not df.empty:
            return df
    return pd.DataFrame(columns=["timestamp", "topic", "difficulty", "count", "faqs"])

def save_history(topic, difficulty, count, faqs, username: str = "guest"):
    df = load_history(username)
    new_row = {
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "topic": topic, "difficulty": difficulty,
        "count": count, "faqs": faqs,
    }
    df = pd.concat([pd.DataFrame([new_row]), df], ignore_index=True)
    df.to_csv(history_file(username), index=False)

# ── Ratings ───────────────────────────────────────────────────────────────────
RATINGS_FILE = "faq_ratings.json"

def load_ratings() -> dict:
    if os.path.exists(RATINGS_FILE):
        with open(RATINGS_FILE, "r") as f:
            return json.load(f)
    return {}

def save_rating(faq_key: str, rating: str, username: str):
    ratings = load_ratings()
    if faq_key not in ratings:
        ratings[faq_key] = {"up": [], "down": []}
    ratings[faq_key]["up"]   = [u for u in ratings[faq_key]["up"]   if u != username]
    ratings[faq_key]["down"] = [u for u in ratings[faq_key]["down"] if u != username]
    ratings[faq_key][rating].append(username)
    with open(RATINGS_FILE, "w") as f:
        json.dump(ratings, f, indent=2)

def get_rating(faq_key: str) -> tuple:
    ratings = load_ratings()
    r = ratings.get(faq_key, {"up": [], "down": []})
    return len(r["up"]), len(r["down"])


# ── Email sender ──────────────────────────────────────────────────────────────
def send_faq_email(sender_email: str, sender_password: str,
                   recipient: str, topic: str,
                   items: list[dict], pdf_bytes: bytes) -> bool:
    """Send FAQ as email with PDF attachment via Gmail SMTP."""
    try:
        msg = MIMEMultipart()
        msg["From"]    = sender_email
        msg["To"]      = recipient
        msg["Subject"] = f"FAQ Forge — {topic} ({len(items)} Questions)"

        # HTML body
        rows = "".join(
            f"<tr style='border-bottom:1px solid #e0e0e0'>"
            f"<td style='padding:10px;font-weight:600;color:#6c4ef0;vertical-align:top'>Q{i}.</td>"
            f"<td style='padding:10px;vertical-align:top'><b>{item['q']}</b><br>"
            f"<span style='color:#555'>{item['a']}</span></td></tr>"
            for i, item in enumerate(items, 1)
        )
        html = f"""
<div style="font-family:Arial,sans-serif;max-width:700px;margin:auto">
  <div style="background:linear-gradient(135deg,#6c4ef0,#06b6d4);padding:24px 28px;border-radius:12px 12px 0 0">
    <h1 style="color:white;margin:0;font-size:22px">◈ FAQ Forge</h1>
    <p style="color:rgba(255,255,255,0.8);margin:4px 0 0;font-size:13px">{topic}</p>
  </div>
  <div style="background:#f9f9fb;padding:20px 28px;border-radius:0 0 12px 12px">
    <table style="width:100%;border-collapse:collapse">{rows}</table>
    <p style="font-size:11px;color:#aaa;margin-top:20px;text-align:center">
      Generated by FAQ Forge · {datetime.now().strftime("%B %d, %Y")}
    </p>
  </div>
</div>"""
        msg.attach(MIMEText(html, "html"))

        # PDF attachment
        part = MIMEBase("application", "octet-stream")
        part.set_payload(pdf_bytes)
        encoders.encode_base64(part)
        slug = topic.lower().replace(" ", "_")
        part.add_header("Content-Disposition", f'attachment; filename="faq_{slug}.pdf"')
        msg.attach(part)

        # Send via Gmail
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, recipient, msg.as_string())
        return True
    except Exception as e:
        raise e

def diff_badge(diff: str) -> str:
    cls = {"Easy": "badge-easy", "Medium": "badge-medium", "Hard": "badge-hard"}.get(diff, "")
    return f'<span class="diff-badge {cls}">{diff}</span>'

def parse_faqs(raw: str) -> list[dict]:
    """Parse Q/A pairs from model output."""
    items = []
    lines = raw.strip().split("\n")
    current_q, current_a = "", ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        # Match Q1:, Q2:, etc. or **Q1:** patterns
        q_match = re.match(r"^\*?\*?Q\d+[.:]\*?\*?\s*(.*)", line, re.IGNORECASE)
        a_match = re.match(r"^\*?\*?A\d+[.:]\*?\*?\s*(.*)", line, re.IGNORECASE)
        if q_match:
            if current_q:
                items.append({"q": current_q, "a": current_a.strip()})
            current_q = q_match.group(1).strip("* ")
            current_a = ""
        elif a_match:
            current_a = a_match.group(1).strip("* ")
        elif current_a:
            current_a += " " + line
    if current_q:
        items.append({"q": current_q, "a": current_a.strip()})
    return items


def remove_duplicates(items: list[dict], threshold: float = 0.6) -> tuple:
    """Remove similar questions using word overlap ratio."""
    def similarity(a: str, b: str) -> float:
        wa = set(a.lower().split())
        wb = set(b.lower().split())
        if not wa or not wb:
            return 0.0
        return len(wa & wb) / len(wa | wb)

    unique, removed = [], []
    for item in items:
        is_dup = any(similarity(item["q"], u["q"]) >= threshold for u in unique)
        if is_dup:
            removed.append(item)
        else:
            unique.append(item)
    return unique, removed

# ─── Export Helpers ───────────────────────────────────────────────────────────

def build_json(items: list[dict], topic: str, difficulty: str, language: str) -> bytes:
    """Return FAQ list as JSON bytes."""
    data = {
        "topic": topic,
        "difficulty": difficulty,
        "language": language,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "faqs": [{"id": i+1, "question": item["q"], "answer": item["a"]} for i, item in enumerate(items)]
    }
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def build_docx(items: list[dict], topic: str, difficulty: str, language: str,
               categories: dict = None) -> bytes:
    """Build a styled Word .docx file."""
    doc = DocxDocument()

    # ── Page margins ──
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    section = doc.sections[0]
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

    # ── Title ──
    title = doc.add_heading(f"FAQ Forge — {topic}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.runs[0]
    run.font.color.rgb = RGBColor(0x6c, 0x4e, 0xf0)
    run.font.size = Pt(26)

    # ── Meta line ──
    meta = doc.add_paragraph()
    meta.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y · %H:%M')}   ·   "
                 f"Difficulty: {difficulty}   ·   Language: {language}   ·   "
                 f"{len(items)} entries").font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(0x7c, 0x76, 0x9e)

    doc.add_paragraph()  # spacer

    # ── FAQ Items ──
    if categories:
        for cat, cat_items in categories.items():
            # Category heading
            ch = doc.add_heading(cat, level=2)
            ch.runs[0].font.color.rgb = RGBColor(0x22, 0xd3, 0xee)
            for item in cat_items:
                _add_faq_item(doc, item)
    else:
        for item in items:
            _add_faq_item(doc, item)

    # ── Footer ──
    doc.add_paragraph()
    footer_p = doc.add_paragraph("Generated by FAQ Forge  ◈")
    footer_p.runs[0].font.size = Pt(8)
    footer_p.runs[0].font.color.rgb = RGBColor(0x4a, 0x45, 0x68)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_faq_item(doc, item: dict):
    """Add a single Q&A block to the Word doc."""
    q_para = doc.add_paragraph()
    q_run  = q_para.add_run(f"Q:  {item['q']}")
    q_run.bold = True
    q_run.font.size = Pt(11)
    q_run.font.color.rgb = RGBColor(0x8a, 0x6e, 0xff)

    a_para = doc.add_paragraph()
    a_run  = a_para.add_run(f"A:  {item['a']}")
    a_run.font.size = Pt(10)
    a_run.font.color.rgb = RGBColor(0x3d, 0x34, 0x60)
    a_para.paragraph_format.left_indent = Inches(0.2)
    a_para.paragraph_format.space_after = Pt(10)

    # Thin separator
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '2')
    bottom.set(qn('w:color'), 'c4bfdf')
    pBdr.append(bottom)
    pPr.append(pBdr)


def categorize_faqs(client, items: list[dict], topic: str) -> dict:
    """Use Gemini to group FAQs into categories."""
    qa_text = "\n".join([f"Q{i+1}: {item['q']}" for i, item in enumerate(items)])
    prompt = f"""Group these FAQ questions about "{topic}" into 2-4 logical categories.
Return ONLY valid JSON in this exact format, nothing else:
{{
  "Category Name 1": [1, 2, 3],
  "Category Name 2": [4, 5]
}}
Where numbers are the Q numbers (1-based index).

Questions:
{qa_text}
"""
    for attempt in range(3):
        try:
            response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            raw = response.text.strip()
            raw = re.sub(r"```json|```", "", raw).strip()
            mapping = json.loads(raw)
            # Build category → items dict
            result = {}
            for cat, indices in mapping.items():
                result[cat] = [items[i-1] for i in indices if 0 < i <= len(items)]
            return result
        except Exception:
            if attempt < 2:
                time.sleep(2)
    return {}


def score_faq_quality(client, items: list[dict], topic: str) -> list[dict]:
    """Use Gemini to score each FAQ answer quality 1-5."""
    qa_text = "\n".join([f"Q{i+1}: {item['q']}\nA{i+1}: {item['a']}" for i, item in enumerate(items)])
    prompt = f"""Rate each FAQ answer below for quality on a scale of 1-5.
1=Poor, 2=Fair, 3=Good, 4=Very Good, 5=Excellent

Return ONLY valid JSON array, nothing else:
[{{"id": 1, "score": 4, "reason": "brief reason"}}, ...]

FAQs about "{topic}":
{qa_text}
"""
    for attempt in range(3):
        try:
            response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            raw = response.text.strip()
            raw = re.sub(r"```json|```", "", raw).strip()
            scores = json.loads(raw)
            # Merge scores back into items
            scored = []
            for i, item in enumerate(items):
                s = next((x for x in scores if x.get("id") == i+1), {})
                scored.append({**item,
                               "score": s.get("score", 0),
                               "reason": s.get("reason", "")})
            return scored
        except Exception:
            if attempt < 2:
                time.sleep(2)
    return items


def build_csv(items: list[dict], topic: str, difficulty: str) -> bytes:
    """Return FAQ list as UTF-8 CSV bytes."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["#", "Topic", "Difficulty", "Question", "Answer"])
    for i, item in enumerate(items, 1):
        writer.writerow([i, topic, difficulty, item["q"], item["a"]])
    return buf.getvalue().encode("utf-8")


def build_pdf(items: list[dict], topic: str, difficulty: str) -> bytes:
    """Return a styled FAQ PDF as bytes using ReportLab."""
    buf = io.BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=20*mm,
        rightMargin=20*mm,
        topMargin=22*mm,
        bottomMargin=22*mm,
    )

    # ── Colour palette ──
    C_BG      = colors.HexColor("#080a0f")
    C_ACCENT  = colors.HexColor("#63c0ff")
    C_TEXT    = colors.HexColor("#e8edf4")
    C_MUTED   = colors.HexColor("#6b7a8f")
    C_SURFACE = colors.HexColor("#0d1117")
    C_BORDER  = colors.HexColor("#1a2d40")
    C_GOLD    = colors.HexColor("#f5c842")
    DIFF_COL  = {"Easy": colors.HexColor("#63ffb4"),
                 "Medium": C_GOLD,
                 "Hard": colors.HexColor("#ff6b6b")}.get(difficulty, C_ACCENT)

    # ── Styles ──
    s_title = ParagraphStyle("title",
        fontName="Helvetica-Bold", fontSize=26,
        textColor=C_TEXT, spaceAfter=4, leading=30)
    s_meta = ParagraphStyle("meta",
        fontName="Helvetica", fontSize=9,
        textColor=C_MUTED, spaceAfter=0, leading=13)
    s_tag = ParagraphStyle("tag",
        fontName="Helvetica-Bold", fontSize=8,
        textColor=DIFF_COL, spaceAfter=0)
    s_q_num = ParagraphStyle("qnum",
        fontName="Helvetica-Bold", fontSize=8,
        textColor=C_ACCENT, spaceAfter=0)
    s_q = ParagraphStyle("q",
        fontName="Helvetica-Bold", fontSize=11,
        textColor=C_TEXT, leading=16, spaceAfter=5)
    s_a = ParagraphStyle("a",
        fontName="Helvetica", fontSize=10,
        textColor=C_MUTED, leading=15, spaceAfter=0)

    story = []

    # ── Header block ──
    story.append(Paragraph("◈  FAQ Forge", s_meta))
    story.append(Spacer(1, 3*mm))
    story.append(Paragraph(topic, s_title))
    story.append(Spacer(1, 1*mm))

    generated_at = datetime.now().strftime("%B %d, %Y · %H:%M")
    story.append(Paragraph(
        f'<font color="#6b7a8f">{generated_at}</font>'
        f'&nbsp;&nbsp;&nbsp;'
        f'<font color="{DIFF_COL.hexval()}">{difficulty}</font>'
        f'&nbsp;&nbsp;&nbsp;'
        f'<font color="#6b7a8f">{len(items)} entries</font>',
        s_meta))
    story.append(Spacer(1, 4*mm))
    story.append(HRFlowable(width="100%", thickness=0.5,
                             color=C_BORDER, spaceAfter=6*mm))

    # ── FAQ items ──
    for i, item in enumerate(items, 1):
        # Q row: number pill + question text side by side
        num_para  = Paragraph(f"Q{i:02d}", s_q_num)
        q_para    = Paragraph(item["q"], s_q)
        a_para    = Paragraph(item["a"], s_a)

        tbl_data = [
            [num_para, q_para],
            ["",        a_para],
        ]
        col_w = [12*mm, doc.width - 12*mm]
        tbl = Table(tbl_data, colWidths=col_w)
        tbl.setStyle(TableStyle([
            ("VALIGN",       (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING",  (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING",   (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING",(0, 0), (-1, -1), 3),
        ]))
        story.append(tbl)

        # thin separator between items (not after last)
        if i < len(items):
            story.append(Spacer(1, 3*mm))
            story.append(HRFlowable(width="100%", thickness=0.3,
                                     color=C_BORDER, spaceAfter=3*mm))

    # ── Footer via canvas override ──
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(C_MUTED)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20*mm, 13*mm, "Generated by FAQ Forge")
        canvas.drawRightString(A4[0] - 20*mm, 13*mm,
                               f"Page {doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ─── File Reading Helpers ─────────────────────────────────────────────────────

def extract_text_from_pdf(uploaded_file) -> str:
    """Extract all text from an uploaded PDF file."""
    reader = PdfReader(uploaded_file)
    pages_text = []
    for i, page in enumerate(reader.pages):
        t = page.extract_text()
        if t and t.strip():
            pages_text.append(f"[Page {i+1}]\n{t.strip()}")
    return "\n\n".join(pages_text)


def describe_image_with_gemini(client, image_bytes: bytes, mime_type: str) -> str:
    """Use Gemini vision to extract text/description from an image."""
    b64 = base64.standard_b64encode(image_bytes).decode()
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[
            {
                "parts": [
                    {
                        "inline_data": {
                            "mime_type": mime_type,
                            "data": b64,
                        }
                    },
                    {
                        "text": (
                            "Extract and describe ALL text, content, diagrams, tables, "
                            "and information visible in this image in detail. "
                            "Format it as clean readable text that can be used as documentation."
                        )
                    },
                ]
            }
        ],
    )
    return response.text


@st.cache_resource
def get_client(api_key: str):
    return genai.Client(api_key=api_key, http_options={"api_version": "v1"})

def generate_faq(client, content: str, topic: str, difficulty: str, n: int, language: str = "English") -> str:
    difficulty_guide = {
        "Easy":   "straightforward, beginner-friendly definitions and basic how-tos",
        "Medium": "practical use-cases, common pitfalls, integration patterns, and step-by-step how-tos",
        "Hard":   "edge cases, performance implications, architectural trade-offs, and advanced troubleshooting",
    }
    prompt = f"""You are a senior technical writer producing polished product documentation.

Generate EXACTLY {n} FAQ entries about "{topic}" using ONLY the content provided.
Difficulty: {difficulty} — focus on {difficulty_guide[difficulty]}.
IMPORTANT: Write ALL questions and answers in {language} language only.

STRICT FORMAT — no preamble, no commentary, nothing else:
Q1: [Question text]
A1: [Answer text]

Q2: [Question text]
A2: [Answer text]

(continue for all {n} entries)

--- DOCUMENTATION ---
{content}
--- END ---
"""
    for attempt in range(3):
        try:
            response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            return response.text
        except Exception as e:
            if attempt < 2:
                time.sleep(3)
            else:
                raise e


def auto_detect_topic(client, content: str) -> str:
    """Use Gemini to auto-detect the main topic from the document."""
    prompt = f"""Read the following document and return ONLY a short topic name (2-5 words max).
No explanation, no punctuation, just the topic words.

Document:
{content[:2000]}
"""
    for attempt in range(3):
        try:
            response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            return response.text.strip()
        except Exception as e:
            if attempt < 2:
                time.sleep(3)
            else:
                raise e


def generate_summary(client, items: list, topic: str, language: str = "English") -> str:
    """Generate a one-paragraph summary of all FAQs."""
    qa_text = "\n".join([f"Q: {i['q']}\nA: {i['a']}" for i in items])
    prompt = f"""Based on these FAQ entries about "{topic}", write a single concise paragraph summary 
(3-4 sentences) that captures the key information covered. Write in {language} language.
No bullet points, just a clean paragraph.

FAQs:
{qa_text}
"""
    for attempt in range(3):
        try:
            response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            return response.text.strip()
        except Exception as e:
            if attempt < 2:
                time.sleep(3)
            else:
                raise e

# ─── API Key ─────────────────────────────────────────────────────────────────
GEMINI_API_KEY = "AIzaSyCXnd1czbPj2pS2M1fws67j55BjtkMK16g"
api_key = GEMINI_API_KEY

# ─── Session State Initialization ────────────────────────────────────────────
if "logged_in"  not in st.session_state:
    st.session_state.logged_in  = False
if "username"   not in st.session_state:
    st.session_state.username   = ""
if "dark_mode"  not in st.session_state:
    st.session_state.dark_mode  = True

# ─── Login Wall ───────────────────────────────────────────────────────────────
if not st.session_state.logged_in:
    st.markdown("""
<div class="login-wrapper">
  <div class="login-logo">◈</div>
  <div class="login-title">FAQ Forge</div>
  <div class="login-sub">Sign in to continue</div>
</div>
""", unsafe_allow_html=True)

    # Center the form
    _, col_c, _ = st.columns([1, 2, 1])
    with col_c:
        login_tab, reg_tab = st.tabs(["🔐 Sign In", "✏️ Register"])

        with login_tab:
            l_user = st.text_input("Username", key="l_user", placeholder="Enter username")
            l_pass = st.text_input("Password", type="password", key="l_pass", placeholder="Enter password")
            if st.button("Sign In", use_container_width=True, key="login_btn"):
                if verify_user(l_user, l_pass):
                    st.session_state.logged_in = True
                    st.session_state.username  = l_user
                    st.rerun()
                else:
                    st.error("❌ Invalid username or password")
            st.markdown("""
<div style="font-size:0.72rem;color:#4a4568;text-align:center;margin-top:0.8rem">
  Default: admin / admin123
</div>""", unsafe_allow_html=True)

        with reg_tab:
            r_user = st.text_input("Username", key="r_user", placeholder="Choose a username")
            r_pass = st.text_input("Password", type="password", key="r_pass", placeholder="Choose a password")
            r_pass2 = st.text_input("Confirm Password", type="password", key="r_pass2", placeholder="Repeat password")
            if st.button("Create Account", use_container_width=True, key="reg_btn"):
                if not r_user or not r_pass:
                    st.warning("Fill all fields")
                elif r_pass != r_pass2:
                    st.error("❌ Passwords don't match")
                elif register_user(r_user, r_pass):
                    st.success(f"✓ Account created! Sign in as **{r_user}**")
                else:
                    st.error("❌ Username already exists")
    st.stop()

# ── Logged-in user ──
current_user = st.session_state.username

# ─── Hide sidebar completely ──────────────────────────────────────────────────
st.markdown("""
<style>
[data-testid="stSidebar"] { display: none !important; }
[data-testid="collapsedControl"] { display: none !important; }
.main .block-container { padding-left: 3.5rem !important; }
</style>
""", unsafe_allow_html=True)

# ─── Header ──────────────────────────────────────────────────────────────────
col_logo, col_user_area = st.columns([6, 2])
with col_logo:
    st.markdown("""
<div class="page-header">
  <div class="header-icon">◈</div>
  <div>
    <div class="header-title">FAQ Forge</div>
  </div>
</div>
""", unsafe_allow_html=True)

with col_user_area:
    st.markdown('<div style="height:1rem"></div>', unsafe_allow_html=True)
    st.markdown(f"""
<div class="user-badge">👤 {current_user}</div>
""", unsafe_allow_html=True)
    uc1, uc2 = st.columns(2)
    with uc1:
        mode_label = "☀️" if st.session_state.dark_mode else "🌙"
        if st.button(mode_label, use_container_width=True):
            st.session_state.dark_mode = not st.session_state.dark_mode
            st.rerun()
    with uc2:
        if st.button("🚪 Logout", use_container_width=True):
            st.session_state.logged_in = False
            st.session_state.username  = ""
            st.rerun()

if not st.session_state.dark_mode:
    st.markdown("""
<script>document.body.classList.add('light-mode');</script>
<style>
.stApp { background-color: #f4f2ff !important; }
.stApp * { color: #1a1035 !important; }
.faq-block, .config-panel, .faq-item { background: #ffffff !important; border-color: rgba(108,78,240,0.15) !important; }
.faq-q { color: #1a1035 !important; }
.faq-a { color: #3d3460 !important; }
.faq-result-title { color: #1a1035 !important; }
h2, h3 { color: #1a1035 !important; }
.history-topic { color: #1a1035 !important; }
</style>
""", unsafe_allow_html=True)

# ─── Main Layout ─────────────────────────────────────────────────────────────
col_doc, col_cfg = st.columns([3, 2], gap="large")

with col_doc:
    tab_text, tab_pdf, tab_img = st.tabs(["✏️  Type / Paste Text", "📄  Upload PDF", "🖼️  Upload Image"])

    extracted_content = ""

    with tab_text:
        text_input = st.text_area(
            "Documentation Content",
            placeholder="Paste product docs, release notes, API references, help articles, or any reference text…",
            height=280,
        )
        extracted_content = text_input

    with tab_pdf:
        pdf_file = st.file_uploader(
            "Upload a PDF document",
            type=["pdf"],
            help="Text will be extracted automatically from all pages",
        )
        if pdf_file:
            with st.spinner("Extracting text from PDF…"):
                try:
                    pdf_text = extract_text_from_pdf(pdf_file)
                    extracted_content = pdf_text
                    st.markdown(f"""
<div style="background:rgba(52,211,153,0.08);border:1px solid rgba(52,211,153,0.22);
     border-left:3px solid #34d399;border-radius:8px;padding:0.6rem 1rem;margin-top:0.8rem;
     font-size:0.75rem;color:#34d399;font-family:'Plus Jakarta Sans',sans-serif;">
  ✓ <strong>{pdf_file.name}</strong> ready — {len(pdf_text):,} characters extracted
</div>
""", unsafe_allow_html=True)
                except Exception as e:
                    st.error(f"Could not read PDF: {e}")

    with tab_img:
        img_file = st.file_uploader(
            "Upload an image (screenshot, diagram, photo of document)",
            type=["png", "jpg", "jpeg", "webp"],
            help="Gemini vision will read and extract all content from the image",
        )
        if img_file:
            col_prev, col_ext = st.columns([1, 1])
            with col_prev:
                st.image(img_file, caption=img_file.name, use_column_width=True)
            with col_ext:
                with st.spinner("Reading image with Gemini Vision…"):
                    try:
                        client_instance = get_client(api_key)
                        img_bytes = img_file.read()
                        mime = f"image/{img_file.type.split('/')[-1]}"
                        img_text = describe_image_with_gemini(client_instance, img_bytes, mime)
                        extracted_content = img_text
                        st.markdown(f"""
<div class="extracted-label">✓ Content extracted from image</div>
<div class="extracted-preview">{img_text[:800]}{"…" if len(img_text) > 800 else ""}</div>
""", unsafe_allow_html=True)
                    except Exception as e:
                        st.error(f"Image reading failed: {e}")

    # Use whichever tab provided content
    text = extracted_content

with col_cfg:
    st.markdown('<div class="config-panel">', unsafe_allow_html=True)

    # ── Auto-detect topic ──
    col_topic, col_auto = st.columns([3, 1])
    with col_topic:
        topic = st.text_input(
            "FAQ Topic",
            placeholder="e.g. Authentication, Webhooks, Pricing",
        )
    with col_auto:
        auto_btn = st.button("🔍 Auto", use_container_width=True, help="Auto-detect topic from document")

    if auto_btn:
        if not text.strip():
            st.warning("📄 Paste or upload content first to auto-detect topic.")
        else:
            with st.spinner("Detecting topic…"):
                try:
                    _client = get_client(api_key)
                    detected = auto_detect_topic(_client, text)
                    st.session_state["detected_topic"] = detected
                    st.success(f"✓ Topic detected: **{detected}**")
                except Exception as e:
                    st.error(f"Detection failed: {e}")

    if "detected_topic" in st.session_state and not topic:
        topic = st.session_state["detected_topic"]

    # ── Language selector ──
    language = st.selectbox(
        "Output Language",
        ["English", "Telugu", "Hindi", "Tamil", "Kannada", "Malayalam",
         "Bengali", "Marathi", "Gujarati", "French", "Spanish", "German", "Arabic", "Chinese"],
        help="FAQs will be generated in this language",
    )

    difficulty = st.selectbox(
        "Difficulty Level",
        ["Easy", "Medium", "Hard"],
        help="Easy = concepts · Medium = how-tos · Hard = edge cases",
    )

    count = st.slider("Number of FAQs", min_value=3, max_value=20, value=5)

    # Difficulty descriptor
    desc = {
        "Easy":   "Basic concepts & definitions for new users",
        "Medium": "Practical how-tos & common integration patterns",
        "Hard":   "Edge cases, trade-offs & advanced troubleshooting",
    }[difficulty]
    st.markdown(f"""
<div style="background:rgba(138,110,255,0.07);border-left:3px solid #8a6eff;
     border-radius:0 8px 8px 0;padding:0.7rem 1rem;margin:0.5rem 0 1rem;
     font-size:0.78rem;color:#7c769e;line-height:1.5">{desc}</div>
""", unsafe_allow_html=True)

    # ── Phase 2 options ──
    st.markdown('<div style="height:0.3rem"></div>', unsafe_allow_html=True)
    col_opt1, col_opt2, col_opt3 = st.columns(3)
    with col_opt1:
        enable_categories = st.checkbox("📂 Categories", value=False,
                                        help="Auto-group FAQs by topic")
    with col_opt2:
        enable_quality    = st.checkbox("⭐ Quality", value=False,
                                        help="Score each answer 1-5")
    with col_opt3:
        enable_dedup      = st.checkbox("🔍 Dedup", value=True,
                                        help="Remove duplicate questions")

    st.markdown('<div style="height:0.5rem"></div>', unsafe_allow_html=True)
    generate_btn = st.button("◈  Generate FAQs", use_container_width=True)
    st.markdown('</div>', unsafe_allow_html=True)

# ─── Generation ──────────────────────────────────────────────────────────────
if generate_btn:
    if not api_key:
        st.error("🔑 API key not configured. Please set GEMINI_API_KEY in the code.")
    elif not text.strip():
        st.warning("📄 Please paste some documentation content first.")
    elif not topic.strip():
        st.warning("🏷️ Please enter a FAQ topic.")
    else:
        # ── Progress bar generation ───────────────────────────────────────────
        progress_bar = st.progress(0, text="Initializing…")
        status_box   = st.empty()
        try:
            t0 = time.time()

            progress_bar.progress(10, text="🔗 Connecting to Gemini…")
            time.sleep(0.3)
            client_instance = get_client(api_key)

            progress_bar.progress(30, text=f"✍️ Generating {count} FAQs in {language}…")
            raw = generate_faq(client_instance, text, topic, difficulty, count, language)
            elapsed = time.time() - t0

            progress_bar.progress(60, text="📋 Parsing results…")
            save_history(topic, difficulty, count, raw, current_user)
            items = parse_faqs(raw)

            # ── Duplicate detector ────────────────────────────────────────────
            progress_bar.progress(70, text="🔍 Checking duplicates…")
            if enable_dedup and items:
                items, dupes = remove_duplicates(items)
                if dupes:
                    status_box.markdown(f"""
<div style="background:rgba(251,191,36,0.07);border:1px solid rgba(251,191,36,0.2);
     border-left:3px solid #fbbf24;border-radius:0 8px 8px 0;
     padding:0.6rem 1rem;margin-bottom:0.8rem;font-size:0.75rem;color:#fbbf24;">
  🔍 Removed <strong>{len(dupes)}</strong> duplicate question(s) automatically
</div>
""", unsafe_allow_html=True)

                # ── Language badge ──
                lang_colors = {
                    "English": "#22d3ee", "Telugu": "#f472b6", "Hindi": "#fb923c",
                    "Tamil": "#a78bfa", "Kannada": "#34d399", "Malayalam": "#fbbf24",
                }
                lang_color = lang_colors.get(language, "#8a6eff")

                st.markdown(f"""
<div class="faq-result-wrapper">
  <div class="faq-result-header">
    <span class="faq-result-title">{topic} — FAQs</span>
    <span style="display:flex;gap:0.5rem;align-items:center;flex-wrap:wrap">
      {diff_badge(difficulty)}
      <span class="faq-tag" style="color:{lang_color};border-color:{lang_color}40;background:{lang_color}12">{language}</span>
      <span class="faq-tag">{len(items) or count} entries</span>
      <span class="faq-tag">{elapsed:.1f}s</span>
    </span>
  </div>
</div>
""", unsafe_allow_html=True)

                if items:
                    # ── FAQ Categories ────────────────────────────────────────
                    categories = {}
                    if enable_categories:
                        progress_bar.progress(78, text="📂 Grouping into categories…")
                        try:
                            categories = categorize_faqs(client_instance, items, topic)
                        except Exception:
                            categories = {}

                    # ── Quality Scorer ────────────────────────────────────────
                    scored_items = items
                    if enable_quality:
                        progress_bar.progress(88, text="⭐ Scoring answer quality…")
                        try:
                            scored_items = score_faq_quality(client_instance, items, topic)
                        except Exception:
                            scored_items = items

                    progress_bar.progress(100, text="✅ Done!")
                    time.sleep(0.4)
                    progress_bar.empty()

                    # ── Render FAQs ───────────────────────────────────────────
                    SCORE_COLORS = {5: "#34d399", 4: "#22d3ee", 3: "#fbbf24", 2: "#fb923c", 1: "#f472b6"}
                    SCORE_LABELS = {5: "Excellent", 4: "Very Good", 3: "Good", 2: "Fair", 1: "Poor"}

                    if categories:
                        for cat_name, cat_items in categories.items():
                            st.markdown(f"""
<div style="font-size:0.7rem;font-weight:700;letter-spacing:0.15em;text-transform:uppercase;
     color:#22d3ee;margin:1.5rem 0 0.5rem;padding-left:0.2rem;">📂 {cat_name}</div>
""", unsafe_allow_html=True)
                            st.markdown('<div class="faq-block">', unsafe_allow_html=True)
                            for item in cat_items:
                                idx = items.index(item) + 1 if item in items else 0
                                s = next((x for x in scored_items if x.get("q") == item["q"]), {})
                                score = s.get("score", 0)
                                sc = SCORE_COLORS.get(score, "")
                                sl = SCORE_LABELS.get(score, "")
                                score_html = f'<span style="font-size:0.65rem;color:{sc};background:{sc}18;border:1px solid {sc}40;border-radius:100px;padding:0.15rem 0.6rem;margin-left:0.5rem">⭐ {sl}</span>' if score and enable_quality else ""
                                st.markdown(f"""
<div class="faq-item">
  <div class="faq-q">
    <span class="faq-q-num">Q{idx:02d}</span>
    {item['q']}{score_html}
  </div>
  <div class="faq-a">{item['a']}</div>
</div>
""", unsafe_allow_html=True)
                            st.markdown('</div>', unsafe_allow_html=True)
                    else:
                        st.markdown('<div class="faq-block">', unsafe_allow_html=True)
                        for i, item in enumerate(scored_items, 1):
                            score = item.get("score", 0)
                            sc = SCORE_COLORS.get(score, "")
                            sl = SCORE_LABELS.get(score, "")
                            score_html = f'<span style="font-size:0.65rem;color:{sc};background:{sc}18;border:1px solid {sc}40;border-radius:100px;padding:0.15rem 0.6rem;margin-left:0.5rem">⭐ {sl}</span>' if score and enable_quality else ""
                            faq_key = hashlib.md5(item["q"].encode()).hexdigest()[:10]
                            ups, downs = get_rating(faq_key)
                            st.markdown(f"""
<div class="faq-item">
  <div class="faq-q">
    <span class="faq-q-num">Q{i:02d}</span>
    {item['q']}{score_html}
  </div>
  <div class="faq-a">{item['a']}</div>
</div>
""", unsafe_allow_html=True)
                            rb1, rb2, rb3 = st.columns([1, 1, 8])
                            with rb1:
                                if st.button(f"👍 {ups}", key=f"up_{faq_key}_{i}",
                                             use_container_width=True):
                                    save_rating(faq_key, "up", current_user)
                                    st.rerun()
                            with rb2:
                                if st.button(f"👎 {downs}", key=f"dn_{faq_key}_{i}",
                                             use_container_width=True):
                                    save_rating(faq_key, "down", current_user)
                                    st.rerun()
                        st.markdown('</div>', unsafe_allow_html=True)
                else:
                    st.markdown(f'<div class="faq-raw">{raw}</div>', unsafe_allow_html=True)

                # ── FAQ Summarizer ──────────────────────────────────────────
                if items:
                    with st.spinner("Generating summary…"):
                        try:
                            summary = generate_summary(client_instance, items, topic, language)
                            st.markdown(f"""
<div style="background:rgba(34,211,238,0.06);border:1px solid rgba(34,211,238,0.18);
     border-left:3px solid #22d3ee;border-radius:0 10px 10px 0;
     padding:1rem 1.2rem;margin-top:1rem;">
  <div style="font-size:0.65rem;font-weight:700;letter-spacing:0.15em;text-transform:uppercase;
       color:#22d3ee;margin-bottom:0.5rem;">📋 FAQ Summary</div>
  <div style="font-size:0.88rem;color:#c4bfdf;line-height:1.75">{summary}</div>
</div>
""", unsafe_allow_html=True)
                        except Exception:
                            pass

                st.success(f"✓ {len(items) or count} FAQs generated in {language} — {elapsed:.1f}s")

                # ── Download buttons ──────────────────────────────────────────
                if items:
                    slug = topic.lower().replace(" ", "_")
                    st.markdown('<div style="height:0.8rem"></div>', unsafe_allow_html=True)

                    dl_col1, dl_col2, dl_col3, dl_col4 = st.columns(4)

                    with dl_col1:
                        st.download_button(
                            label="⬇ CSV",
                            data=build_csv(items, topic, difficulty),
                            file_name=f"faq_{slug}.csv",
                            mime="text/csv",
                            use_container_width=True,
                        )
                    with dl_col2:
                        st.download_button(
                            label="⬇ PDF",
                            data=build_pdf(items, topic, difficulty),
                            file_name=f"faq_{slug}.pdf",
                            mime="application/pdf",
                            use_container_width=True,
                        )
                    with dl_col3:
                        st.download_button(
                            label="⬇ JSON",
                            data=build_json(items, topic, difficulty, language),
                            file_name=f"faq_{slug}.json",
                            mime="application/json",
                            use_container_width=True,
                        )
                    with dl_col4:
                        if DOCX_AVAILABLE:
                            st.download_button(
                                label="⬇ Word",
                                data=build_docx(items, topic, difficulty, language, categories if enable_categories else None),
                                file_name=f"faq_{slug}.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True,
                            )
                        else:
                            st.button("⬇ Word\n(install python-docx)", disabled=True, use_container_width=True)

                # ── Email FAQ ─────────────────────────────────────────────────
                if items:
                    with st.expander("📧 Email this FAQ"):
                        st.markdown("""
<div style="font-size:0.72rem;color:#7c769e;margin-bottom:0.8rem">
  Sends FAQ as a styled HTML email + PDF attachment via Gmail.
  Use a <a href="https://support.google.com/accounts/answer/185833" target="_blank" style="color:#8a6eff">Gmail App Password</a>.
</div>
""", unsafe_allow_html=True)
                        em_col1, em_col2 = st.columns(2)
                        with em_col1:
                            sender_email = st.text_input("Your Gmail", placeholder="you@gmail.com", key="s_email")
                            sender_pass  = st.text_input("App Password", type="password",
                                                          placeholder="xxxx xxxx xxxx xxxx", key="s_pass")
                        with em_col2:
                            recipient    = st.text_input("Send To", placeholder="recipient@email.com", key="r_email")
                            st.markdown('<div style="height:1.5rem"></div>', unsafe_allow_html=True)
                            send_btn = st.button("📤 Send Email", use_container_width=True, key="send_email")

                        if send_btn:
                            if not sender_email or not sender_pass or not recipient:
                                st.warning("Fill all email fields.")
                            elif "@gmail.com" not in sender_email:
                                st.error("❌ Only Gmail supported. Use a @gmail.com address.")
                            else:
                                with st.spinner("Sending email…"):
                                    try:
                                        pdf_attach = build_pdf(items, topic, difficulty)
                                        send_faq_email(sender_email, sender_pass,
                                                       recipient, topic, items, pdf_attach)
                                        st.success(f"✅ Email sent to **{recipient}** with PDF attached!")
                                    except smtplib.SMTPAuthenticationError:
                                        st.error("❌ Authentication failed! Make sure you are using a **Gmail App Password**, not your regular password. Generate one at: myaccount.google.com/apppasswords")
                                    except smtplib.SMTPRecipientsRefused:
                                        st.error("❌ Recipient email address is invalid.")
                                    except smtplib.SMTPException as e:
                                        st.error(f"❌ SMTP Error: {e}")
                                    except Exception as e:
                                        st.error(f"❌ Failed: {e}")

        except Exception as e:
            progress_bar.empty()
            st.error(f"Generation failed: {e}")

# ─── History & Compare ───────────────────────────────────────────────────────
st.markdown("---")
hist_tab, compare_tab = st.tabs(["🗂️ My History", "⚖️ Compare FAQs"])

with hist_tab:
    hist_df = load_history(current_user)
    col_h1, col_h2 = st.columns([5, 1])
    with col_h1:
        st.markdown(f"### {current_user}'s History")
    with col_h2:
        if not hist_df.empty:
            st.markdown('<div class="clear-btn">', unsafe_allow_html=True)
            if st.button("Clear all", use_container_width=True):
                os.remove(history_file(current_user))
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    if hist_df.empty:
        st.markdown("""
<div style="background:#16161f;border:1px dashed rgba(138,110,255,0.18);border-radius:12px;
     padding:2.5rem;text-align:center;color:#4a4568;
     font-family:'Plus Jakarta Sans',sans-serif;font-size:0.88rem;">
  No FAQs generated yet — run your first generation above
</div>
""", unsafe_allow_html=True)
    else:
        for _, row in hist_df.iterrows():
            topic_val = row.get("topic", "—")
            diff_val  = row.get("difficulty", "")
            count_val = row.get("count", "?")
            ts_val    = row.get("timestamp", "")
            st.markdown(f"""
<div class="history-row">
  <div>
    <div class="history-topic">{topic_val}</div>
    <div class="history-meta">{ts_val} &nbsp;·&nbsp; {count_val} entries</div>
  </div>
  {diff_badge(diff_val)}
</div>
""", unsafe_allow_html=True)

        with st.expander("📄 View full FAQ text"):
            for _, row in hist_df.iterrows():
                st.markdown(f"""
<div style="margin-bottom:0.5rem;display:flex;align-items:center;gap:0.7rem">
  <span style="font-family:'Plus Jakarta Sans',sans-serif;font-weight:700;font-size:0.95rem;color:#f0eeff">{row.get('topic','')}</span>
  {diff_badge(row.get('difficulty',''))}
  <span style="font-size:0.68rem;color:#4a4568;font-family:'JetBrains Mono',monospace">{row.get('timestamp','')}</span>
</div>
""", unsafe_allow_html=True)
                items_h = parse_faqs(str(row.get("faqs", "")))
                if items_h:
                    st.markdown('<div class="faq-block" style="margin-bottom:0.8rem">', unsafe_allow_html=True)
                    for i, item in enumerate(items_h, 1):
                        st.markdown(f"""
<div class="faq-item">
  <div class="faq-q"><span class="faq-q-num">Q{i:02d}</span>{item['q']}</div>
  <div class="faq-a">{item['a']}</div>
</div>""", unsafe_allow_html=True)
                    st.markdown('</div>', unsafe_allow_html=True)
                    slug = str(row.get("topic","faq")).lower().replace(" ","_")
                    diff = str(row.get("difficulty",""))
                    hcol1, hcol2, hcol3 = st.columns([1,1,3])
                    with hcol1:
                        st.download_button("⬇ CSV",
                            build_csv(items_h, str(row.get("topic","")), diff),
                            file_name=f"faq_{slug}.csv", mime="text/csv",
                            use_container_width=True, key=f"csv_{slug}_{_}")
                    with hcol2:
                        st.download_button("⬇ PDF",
                            build_pdf(items_h, str(row.get("topic","")), diff),
                            file_name=f"faq_{slug}.pdf", mime="application/pdf",
                            use_container_width=True, key=f"pdf_{slug}_{_}")
                else:
                    st.markdown(f'<div class="faq-raw">{row.get("faqs","")}</div>', unsafe_allow_html=True)
                st.write("")

with compare_tab:
    st.markdown("### Compare Two FAQ Sets")
    hist_df2 = load_history(current_user)
    if hist_df2.empty or len(hist_df2) < 2:
        st.info("Generate at least 2 FAQ sets to compare them.")
    else:
        topics_list = hist_df2["topic"].tolist()
        cc1, cc2 = st.columns(2)
        with cc1:
            sel_a = st.selectbox("FAQ Set A", range(len(topics_list)),
                                  format_func=lambda i: f"{topics_list[i]} ({hist_df2.iloc[i]['timestamp']})",
                                  key="cmp_a")
        with cc2:
            sel_b = st.selectbox("FAQ Set B", range(len(topics_list)),
                                  format_func=lambda i: f"{topics_list[i]} ({hist_df2.iloc[i]['timestamp']})",
                                  index=min(1, len(topics_list)-1),
                                  key="cmp_b")

        if sel_a != sel_b:
            items_a = parse_faqs(str(hist_df2.iloc[sel_a]["faqs"]))
            items_b = parse_faqs(str(hist_df2.iloc[sel_b]["faqs"]))
            ca, cb = st.columns(2)
            with ca:
                st.markdown(f"""
<div style="font-size:0.7rem;font-weight:700;letter-spacing:0.12em;text-transform:uppercase;
     color:#8a6eff;margin-bottom:0.5rem">Set A — {topics_list[sel_a]}</div>
""", unsafe_allow_html=True)
                st.markdown('<div class="faq-block">', unsafe_allow_html=True)
                for i, item in enumerate(items_a, 1):
                    st.markdown(f"""
<div class="faq-item">
  <div class="faq-q"><span class="faq-q-num">Q{i:02d}</span>{item['q']}</div>
  <div class="faq-a">{item['a']}</div>
</div>""", unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)
            with cb:
                st.markdown(f"""
<div style="font-size:0.7rem;font-weight:700;letter-spacing:0.12em;text-transform:uppercase;
     color:#22d3ee;margin-bottom:0.5rem">Set B — {topics_list[sel_b]}</div>
""", unsafe_allow_html=True)
                st.markdown('<div class="faq-block">', unsafe_allow_html=True)
                for i, item in enumerate(items_b, 1):
                    st.markdown(f"""
<div class="faq-item">
  <div class="faq-q"><span class="faq-q-num">Q{i:02d}</span>{item['q']}</div>
  <div class="faq-a">{item['a']}</div>
</div>""", unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)
        else:
            st.warning("Select two different FAQ sets to compare.")

# ─── Footer ──────────────────────────────────────────────────────────────────
st.markdown("---")
st.markdown(f"""
<div style="display:flex;justify-content:space-between;align-items:center;
     font-size:0.72rem;color:#4a4568;padding:0.5rem 0 1rem">
  <span>◈ FAQ Forge &nbsp;·&nbsp; Logged in as <strong style="color:#8a6eff">{current_user}</strong></span>
  <span>Powered by Gemini 2.5 Flash</span>
</div>
""", unsafe_allow_html=True)

# ─── Auto-generate requirements.txt ──────────────────────────────────────────
REQUIREMENTS = """streamlit>=1.32.0
google-genai>=0.5.0
pypdf>=4.0.0
reportlab>=4.0.0
python-docx>=1.1.0
pandas>=2.0.0
"""
if not os.path.exists("requirements.txt"):
    with open("requirements.txt", "w") as f:
        f.write(REQUIREMENTS)